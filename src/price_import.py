"""Импорт прайс-листа из Excel-файла бухгалтерии заказчика.

Руководитель выгружает прайс из 1С/МойСклад в Excel и загружает файл
на страницу прайса. Здесь мы читаем таблицу: находим материалы, работы,
наценку и доставку и сравниваем с действующим прайсом, чтобы страница
показала, что изменится, ещё до подтверждения.
"""
from openpyxl import load_workbook

# Какими словами в выгрузках бухгалтерии начинаются нужные заголовки
NAME_WORDS = ("материал", "наименование", "номенклат", "товар", "продукц",
              "позиция", "работа", "услуг")
PRICE_WORDS = ("цена", "стоимость", "тариф")
UNIT_WORDS = ("ед", "изм")
WORK_WORDS = ("работ", "операц", "услуг")
MARGIN_WORDS = ("нацен",)
DELIVERY_WORDS = ("доставк",)
# Строки отчётов и итогов в прайс не идут
SKIP_WORDS = ("итог", "исполь", "скидк", "нацен", "доставк")


def _text(value) -> str:
    """Ячейку любого типа приводим к строке без пробелов по краям."""
    if value is None:
        return ""
    return str(value).strip()


def _number(value):
    """Цена из ячейки: числовой тип или строка с запятой; иначе None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    digits = _text(value).replace(" ", "").replace(",", ".")
    try:
        return round(float(digits), 2)
    except ValueError:
        return None


def _money(value: float) -> str:
    """Число с пробелами разрядов: 12500.0 превращается в 12 500."""
    return f"{value:,.0f}".replace(",", " ")


def _cell(row, col):
    """Значение ячейки по колонке; за краем строки — None."""
    if col is None or col >= len(row):
        return None
    return row[col]


def _column_index(texts, words):
    """Номер колонки по ключевому слову заголовка (первая слева)."""
    for index, text in enumerate(texts):
        for word in words:
            if word in text:
                return index
    return None


def _sections(rows) -> list:
    """Строки-заголовки таблиц с ценами: материалы и работы отдельно."""
    found = []
    for row_index, row in enumerate(rows):
        texts = [_text(cell).lower() for cell in row]
        name_col = _column_index(texts, NAME_WORDS)
        price_col = _column_index(texts, PRICE_WORDS)
        if name_col is None or price_col is None:
            continue
        is_work = any(any(word in text for word in WORK_WORDS)
                      for text in texts)
        found.append({
            "row": row_index,
            "name": name_col,
            "price": price_col,
            "unit": _column_index(texts, UNIT_WORDS),
            "works": is_work,
        })
    return found


def _collect_materials(rows, section, stop, store: dict) -> None:
    """Материалы из строк под заголовком таблицы материалов."""
    for row in rows[section["row"] + 1:stop]:
        name = _text(_cell(row, section["name"]))
        if not name:
            continue
        low = name.lower()
        if any(word in low for word in SKIP_WORDS):
            continue
        price = _number(_cell(row, section["price"]))
        if price is None:
            continue
        unit = _text(_cell(row, section["unit"])) or "ед."
        store[name] = {"price_per_unit": round(price, 2), "unit": unit}


def _collect_rates(rows, section, stop, store: dict) -> None:
    """Работы из строк под заголовком таблицы работ."""
    for row in rows[section["row"] + 1:stop]:
        name = _text(_cell(row, section["name"]))
        if not name:
            continue
        low = name.lower()
        if any(word in low for word in SKIP_WORDS):
            continue
        price = _number(_cell(row, section["price"]))
        if price is not None:
            store[name] = round(price, 2)


def _find_amount(rows, words):
    """Число возле надписи («Наценка…», «Доставка…»): справа или снизу."""
    for row_number, row in enumerate(rows):
        texts = [_text(cell).lower() for cell in row]
        for index, text in enumerate(texts):
            if not any(word in text for word in words):
                continue
            for value in row[index + 1:]:
                number = _number(value)
                if number is not None:
                    return number
            for below in rows[row_number + 1:row_number + 4]:
                number = _number(_cell(below, index))
                if number is not None:
                    return number
    return None


def _margin_from(value):
    """Наценка в коэффициент: 25, 0.25 и 1.25 — это одинаковые 25 %.

    Чёткого стандарта нет: бухгалтерия пишет процент (25), Excel
    превращает его в долю (0.25), либо пишут готовый коэффициент (1.25).
    """
    if value is None:
        return None
    if value < 1:
        return round(1 + value, 4)
    if value < 2:
        return round(value, 4)
    return round(1 + value / 100, 4)


def parse_price(path) -> dict:
    """Читаем Excel: материалы, работы, наценка и доставка (если найдены).

    Наценка и доставка могут отсутствовать в выгрузке — тогда None,
    и при загрузке сохранится действующее значение.
    """
    book = load_workbook(path, read_only=True, data_only=True)
    rows = [list(row) for row in book.active.iter_rows(values_only=True)]
    book.close()
    sections = _sections(rows)
    materials: dict = {}
    rates: dict = {}
    for number, section in enumerate(sections):
        stop = (sections[number + 1]["row"] if number + 1 < len(sections)
                else len(rows))
        if section["works"]:
            _collect_rates(rows, section, stop, rates)
        else:
            _collect_materials(rows, section, stop, materials)
    return {
        "materials": materials,
        "rates": rates,
        "margin": _margin_from(_find_amount(rows, MARGIN_WORDS)),
        "delivery": _find_amount(rows, DELIVERY_WORDS),
    }


def _material_text(info: dict) -> str:
    """Цена материала для предпросмотра: 850 ₽/ед."""
    return f"{_money(info['price_per_unit'])} ₽/{info['unit']}"


def _rate_text(value) -> str:
    """Цена работы для предпросмотра: 450 ₽/шт."""
    return f"{_money(value)} ₽/шт."


def _percent_text(margin) -> str:
    """Коэффициент наценки в проценты: 1.25 → 25 %."""
    return f"{round((margin - 1) * 100, 1):g} %"


def diff(current: dict, imported: dict) -> dict:
    """Сравниваем выгрузку с действующим прайсом до подтверждения."""
    rows: list = []
    counts = {"новый": 0, "изменится": 0, "уберётся": 0}
    for name, info in current["materials"].items():
        if name not in imported["materials"]:
            rows.append({"section": "Материал", "name": name,
                         "status": "уберётся", "old": _material_text(info),
                         "new": "—"})
            counts["уберётся"] += 1
    for name, info in imported["materials"].items():
        old = current["materials"].get(name)
        new = _material_text(info)
        if old is None:
            status = "новый"
        elif _material_text(old) == new:
            status = "без изменений"
        else:
            status = "изменится"
            counts["изменится"] += 1
        if status == "новый":
            counts["новый"] += 1
        rows.append({"section": "Материал", "name": name, "status": status,
                     "old": _material_text(old) if old else "—", "new": new})
    for name, value in current["rates"].items():
        if name not in imported["rates"]:
            rows.append({"section": "Работа", "name": name,
                         "status": "уберётся", "old": _rate_text(value),
                         "new": "—"})
            counts["уберётся"] += 1
    for name, value in imported["rates"].items():
        old = current["rates"].get(name)
        new = _rate_text(value)
        if old is None:
            status = "новый"
        elif round(float(old), 2) == value:
            status = "без изменений"
        else:
            status = "изменится"
            counts["изменится"] += 1
        if status == "новый":
            counts["новый"] += 1
        rows.append({"section": "Работа", "name": name, "status": status,
                     "old": _rate_text(old) if old is not None else "—",
                     "new": new})
    margin_new = imported["margin"] if imported["margin"] is not None \
        else current["margin"]
    delivery_new = imported["delivery"] if imported["delivery"] is not None \
        else current["delivery"]
    rows.append({
        "section": "Итог", "name": "Наценка",
        "status": "изменится" if margin_new != current["margin"]
        else "без изменений",
        "old": _percent_text(current["margin"]),
        "new": _percent_text(margin_new),
    })
    rows.append({
        "section": "Итог", "name": "Доставка",
        "status": "изменится" if delivery_new != current["delivery"]
        else "без изменений",
        "old": f"{_money(current['delivery'])} ₽",
        "new": f"{_money(delivery_new)} ₽",
    })
    has_changes = (bool(counts["новый"]) or bool(counts["изменится"])
                   or bool(counts["уберётся"])
                   or margin_new != current["margin"]
                   or delivery_new != current["delivery"])
    return {"rows": rows, "counts": counts, "has_changes": has_changes}

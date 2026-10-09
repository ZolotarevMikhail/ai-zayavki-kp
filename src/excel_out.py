"""Excel-версия таблицы КП: один лист, читаемые колонки и итоги."""
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from src.config import KP_DIR

HEADERS = ["№", "Наименование", "Материал", "Размер", "Кол-во",
           "Работы", "Цена позиции, ₽"]


def build(order: dict):
    """Собираем файл data/kp/KP_<номер>.xlsx и возвращаем путь к нему."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "КП"
    sheet.append([f"Коммерческое предложение № {order['id']}"])
    sheet.append([f"от {order['created_at'][:10]}"])
    sheet.append([])
    sheet.append(HEADERS)
    calc = order["calc"]
    for item in calc["items"]:
        sheet.append([
            item["number"],
            item["name"],
            item["material"],
            item["size"],
            item["count"],
            ", ".join(item.get("operations", [])),
            item["subtotal"],
        ])
    subtotal_all = calc["total"] - calc["delivery"]
    sheet.append([])
    sheet.append(["", "", "", "", "", "Стоимость работ", subtotal_all])
    sheet.append(["", "", "", "", "", "Доставка", calc["delivery"]])
    sheet.append(["", "", "", "", "", "ИТОГО", calc["total"]])
    widths = (5, 32, 14, 16, 8, 24, 16)
    for number, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(number)].width = width
    path = KP_DIR / f"KP_{order['id']}.xlsx"
    workbook.save(path)
    return path

"""Сборка коммерческого предложения по редактируемому Word-шаблону.

Шаблон лежит в data/templates/kp_template.docx и создаётся при первом запуске.
Клиент может открыть его в Word, поменять фирменный вид, логотип, шрифты —
код подставит данные в те же плейсхолдеры: {{ org_name }}, {{ date }},
{{ currency }}, {{ delivery_text }}, {{ total_text }}, {{ validity }}.
В таблице строки позиций повторяются циклом {%tr for item in items %}.
"""
import os

from docx import Document
from docx.shared import Pt
from docxtpl import DocxTemplate

from src.config import KP_DIR, TEMPLATE_PATH
from src.settings import load_settings


def build_docx(order: dict) -> str:
    """Формируем КП и возвращаем путь к файлу."""
    _ensure_template()
    settings = load_settings()
    doc = DocxTemplate(str(TEMPLATE_PATH))
    context = {
        "org_name": settings["org_name"],
        "org_phone": settings["org_phone"],
        "org_email": settings["org_email"],
        "date": order["created_at"][:10],
        "validity": settings["kp_validity"],
        "currency": settings["currency"],
        "items": _table_rows(order),
        "delivery_text": _money(order["calc"]["delivery"]),
        "total_text": _money(order["calc"]["total"]),
        "lead_note": settings["lead_note"],
    }
    doc.render(context)
    path = os.path.join(KP_DIR, f"KP_{order['id']}.docx")
    doc.save(path)
    return path


def _table_rows(order: dict) -> list:
    """Расчёт превращаем в плоские словари для строк таблицы шаблона."""
    rows = []
    for item in order["calc"]["items"]:
        rows.append({
            "number": item["number"],
            "name": item["name"],
            "material": item["material"],
            "size": item["size"],
            "count": item["count"],
            "price_text": _money(item["material_cost"]),
            "subtotal_text": _money(item["subtotal"]),
            "operations": ", ".join(item.get("operations", [])),
        })
    return rows


def _money(amount: float) -> str:
    """Числа как в бухгалтерии: 12 345, без дробных копеек."""
    return f"{amount:,.0f}".replace(",", " ")


def _ensure_template() -> None:
    """Первый запуск: создаём шаблон КП с плейсхолдерами docxtpl."""
    if TEMPLATE_PATH.exists():
        return
    document = Document()

    header = document.add_paragraph()
    header_run = header.add_run("{{ org_name }}")
    header_run.font.bold = True
    header_run.font.size = Pt(16)
    contacts = document.add_paragraph()
    contacts_run = contacts.add_run("{{ org_phone }}   {{ org_email }}")
    contacts_run.font.size = Pt(9)

    title = document.add_heading("Коммерческое предложение", level=1)
    title.alignment = 1
    document.add_paragraph(
        "Дата: {{ date }}. Предложение действует: {{ validity }}."
    )

    # 4 строки: шапка / маркер начала цикла / строка с полями / маркер конца
    table = document.add_table(rows=4, cols=6)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, _HEADERS):
        cell.text = text
        for paragraph in cell.paragraphs:
            for bold_run in paragraph.runs:
                bold_run.font.bold = True
                bold_run.font.size = Pt(10)

    table.rows[1].cells[0].text = "{%tr for item in items %}"
    data_cells = table.rows[2].cells
    data_cells[0].text = "{{ item.number }}"
    data_cells[1].text = "{{ item.name }}"
    data_cells[2].text = "{{ item.material }} {{ item.size }}"
    data_cells[3].text = "{{ item.count }}"
    data_cells[4].text = "{{ item.price_text }} {{ currency }}"
    data_cells[5].text = "{{ item.subtotal_text }} {{ currency }}"
    table.rows[3].cells[0].text = "{%tr endfor %}"

    document.add_paragraph()
    delivery = document.add_paragraph()
    delivery_run = delivery.add_run(
        "Доставка: {{ delivery_text }} {{ currency }}"
    )
    delivery_run.font.size = Pt(11)
    grand = document.add_paragraph()
    grand_run = grand.add_run("ИТОГО: {{ total_text }} {{ currency }}")
    grand_run.font.bold = True
    grand_run.font.size = Pt(14)
    document.add_paragraph("{{ lead_note }}")

    document.save(str(TEMPLATE_PATH))


_HEADERS = (
    "№", "Наименование", "Материал / размер",
    "Кол-во", "Цена, ₽", "Сумма, ₽",
)
